#!/usr/bin/env python3
"""Port selected layouts from The Hoffman Agency's AMEA PowerPoint template onto TI's master.

    python3 tools/port_hoffman_layouts.py <hoffman-template.pptx> [ti-template.pptx] [out.pptx]

Run after tools/add_layouts.py. Takes one representative of each Hoffman layout family
(the colour repeats are skipped), scales it from 13.33 to 10 in, keeps everything above
TI's footer band, swaps Poppins/Aptos for the theme's Arial, and remaps Hoffman's colours
by role into TI's palette: navy and purple panels become teal, lime and light panels
become grey-100, lime text on a dark panel becomes white, titles inherit TI red. Content
pages stay white; the one teal surface is the Key point layout. Hoffman's corner mark is
dropped. Safe to re-run: layouts already present by name are skipped.
"""
import copy
import os
import random
import re
import sys
import uuid

from lxml import etree
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import add_layouts as al  # noqa: E402

HA_SRC = sys.argv[1]
TPL = sys.argv[2] if len(sys.argv) > 2 else "templates/powerpoint/TI_Deck_Template_v1.pptx"
OUT = sys.argv[3] if len(sys.argv) > 3 else TPL

E = 914400
HA_W, HA_H = 13.333, 7.5
SCALE_X = 10 / HA_W
# Hoffman's content band (title top 0.56 in, content bottom 6.98 in) maps onto TI's (0.12 to 4.80 in),
# so nothing reaches TI's footer rule (5.09 in), signature or page number.
Y_A = (4.80 - 0.12) / (6.98 - 0.56)
Y_B = 0.12 - 0.56 * Y_A
BLEED_BOTTOM = 4.85  # anything that ran to Hoffman's bottom edge now stops above TI's footer band
SIZE_STEPS = (9, 10, 11, 12, 14, 16, 18, 20, 24, 28, 32, 36, 44, 54, 60, 72, 90)

# (Hoffman layout, TI name, surface) — surface=True becomes a full teal page with no signature
PORTS = [
    ("HA-1 Title-2 Columns-1", "Statement + two panels", False),
    ("HA-Title+Content-14", "Title + three points", False),
    ("HA-Highlight-7", "Key point · teal", True),
    ("HA-Title+Content-4", "Four columns + photos", False),
    ("HA-Title+Content-3", "Three photos + text · L49", False),
    ("HA-Title+Content-5", "Two columns + photos", False),
    ("HA-CampaignAtAGlance-1", "Programme at a glance · L10", False),
    ("HA-Title+Content-7", "Text + image half", False),
    ("HA-Title+Content-9", "Text + image third", False),
    ("HA-Title+Content-10", "Image two-thirds + text", False),
    ("HA-Title+Content-2", "Title + body + image band", False),
    ("HA-Title+Content-12", "Text + image + two stats", False),
    ("HA-Title+Content-15", "Image + three points · L48", False),
    ("HA-Collage-1", "Photo collage · L42", False),
    ("HA-Team-2", "Team of 10 · L15", False),
    ("HA-Team-Bio-1", "Single bio · L16", False),
]

ORDER = [
    "Cover · L01", "Section divider · L02", "Appendix divider · L02 teal", "Agenda · L03",
    "Title + body", "Two columns", "Statement + two panels", "Title + three points", "Statement · L08",
    "Key point · teal", "Three columns · L12", "Four columns + photos", "Three photos + text · L49",
    "Two columns + photos", "Stat strip · L19", "Chart + takeaway", "Table · L24",
    "Programme at a glance · L10", "Timeline · L20", "Image + text · L28", "Text + image half",
    "Text + image third", "Image two-thirds + text", "Title + body + image band", "Text + image + two stats",
    "Image + three points · L48", "Photo collage · L42", "Team · L15", "Team of 10 · L15", "Single bio · L16",
    "Pull quote · L29", "Title only", "Closing · L30", "Blank",
]

# Hoffman role -> TI colour, by where the colour is used
NAVY = {"fill": "accent3", "text": "tx1", "line": "accent3"}
ROLE_MAP = {
    "tx1": {"fill": "tx1", "text": "tx1", "line": "tx1"}, "dk1": {"fill": "tx1", "text": "tx1", "line": "tx1"},
    "bg1": {"fill": "bg1", "text": "bg1", "line": "bg1"}, "lt1": {"fill": "bg1", "text": "bg1", "line": "bg1"},
    "tx2": NAVY, "dk2": NAVY, "accent3": NAVY,
    "bg2": {"fill": "#F7F7F7", "text": "tx1", "line": "#CCCCCC"}, "lt2": {"fill": "#F7F7F7", "text": "tx1", "line": "#CCCCCC"},
    "accent1": {"fill": "accent3", "text": "accent3", "line": "accent3"},   # purple -> teal
    "accent2": {"fill": "accent5", "text": "accent3", "line": "accent5"},   # lilac -> teal 300
    "accent4": {"fill": "#F7F7F7", "text": "bg1", "line": "#CCCCCC"},       # lime -> grey panel / white text
    "accent5": {"fill": "accent3", "text": "accent3", "line": "accent3"},   # deep teal -> teal
    "accent6": {"fill": "accent5", "text": "accent3", "line": "accent5"},   # cyan -> teal 300
}
SRGB_ROLE = {"D2EB00": "accent4", "145F7B": "accent5", "EEEEEE": "bg2", "182C43": "tx2", "6103B9": "accent1"}
TEXT_TAGS = {qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr"), qn("a:buClr"), qn("a:fontRef")}


def ymap(v):
    return Y_A * v + Y_B


def snap(pt):
    return min(SIZE_STEPS, key=lambda s: (abs(s - pt), s))


def context(node):
    for anc in node.iterancestors():
        if anc.tag in TEXT_TAGS:
            return "text"
        if anc.tag in (qn("a:ln"), qn("a:lnRef")):
            return "line"
    return "fill"


def remap_colours(el, surface):
    for node in list(el.iter(qn("a:schemeClr"), qn("a:srgbClr"))):
        if node.tag == qn("a:srgbClr"):
            role = SRGB_ROLE.get(node.get("val", "").upper())
            if not role:
                continue
        else:
            role = node.get("val")
        target = ROLE_MAP.get(role, {}).get(context(node))
        if surface and context(node) == "text":
            target = "bg1"  # everything reads white on the teal page
        if not target:
            continue
        new = etree.SubElement(node.getparent(), qn("a:srgbClr") if target.startswith("#") else qn("a:schemeClr"))
        new.set("val", target.lstrip("#"))
        for child in node:  # keep tints and transparency
            new.append(copy.deepcopy(child))
        node.addprevious(new)
        node.getparent().remove(node)


def normalise_text(el):
    """Theme font (Arial) everywhere, sizes scaled to the 10 in page and snapped to TI's tiers."""
    for tf in list(el.iter(qn("a:latin"), qn("a:ea"), qn("a:cs"))):
        face = tf.get("typeface", "")
        if face.startswith("+"):
            continue
        parent = tf.getparent()
        if re.search(r"Bold|SemiBold|Black|Heavy", face) and parent.get("b") is None:
            parent.set("b", "1")
        parent.remove(tf)
    for node in el.iter(qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")):
        if node.get("sz"):
            node.set("sz", str(snap(int(node.get("sz")) / 100 * SCALE_X) * 100))
    for node in el.iter():
        for attr in ("marL", "indent", "lIns", "rIns", "tIns", "bIns"):
            if node.get(attr) and node.tag != qn("p:ph"):
                node.set(attr, str(int(int(node.get(attr)) * SCALE_X)))
    for node in el.iter(qn("a:spcPts")):
        node.set("val", str(int(int(node.get("val")) * SCALE_X)))


def place(el):
    xfrm = el.find(qn("p:spPr") + "/" + qn("a:xfrm"))
    if xfrm is None:
        xfrm = el.find(qn("p:xfrm"))
    if xfrm is None:
        return
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    x, y = int(off.get("x")) / E, int(off.get("y")) / E
    w, h = int(ext.get("cx")) / E, int(ext.get("cy")) / E
    x1, x2 = x * SCALE_X, (x + w) * SCALE_X
    y1 = 0.0 if y < 0.05 else ymap(y)
    y2 = BLEED_BOTTOM if y + h > HA_H - 0.05 else ymap(y + h)
    off.set("x", str(int(round(x1 * E)))); off.set("y", str(int(round(y1 * E))))
    ext.set("cx", str(int(round((x2 - x1) * E)))); ext.set("cy", str(int(round((y2 - y1) * E))))


def is_title(el):
    ph = el.find(".//" + qn("p:ph"))
    return ph is not None and ph.get("type") in ("title", "ctrTitle")


def geometry(el):
    xfrm = el.find(qn("p:spPr") + "/" + qn("a:xfrm"))
    if xfrm is None:
        xfrm = el.find(qn("p:xfrm"))
    if xfrm is None:
        return None
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    return tuple(int(v) / E for v in (off.get("x"), off.get("y"), ext.get("cx"), ext.get("cy")))


def port(ti, ha_layout, name, surface, base):
    L = al.Layout(al.clone_layout(ti, base, name))
    L.drop_title()
    if surface:
        L.drop_picture()
        L.drop(lambda s: s.find(".//" + qn("p:cNvPr")).get("name") == "Slide number")
        L.background("accent3")
    src = ha_layout._element.find(qn("p:cSld")).find(qn("p:spTree"))
    for shape in src:
        if shape.tag not in (qn("p:sp"), qn("p:cxnSp"), qn("p:graphicFrame")):
            continue  # groups and pictures are not part of the ported set
        is_ph = shape.find(".//" + qn("p:ph")) is not None
        g = geometry(shape)
        if not is_ph and g and g[0] > 12.4 and g[1] > 6.9 and g[2] < 0.6:
            continue  # Hoffman's corner mark
        if not is_ph and g and g[0] < 0.05 and g[1] < 0.05 and g[2] > 13 and g[3] > 7.4:
            continue  # full-page colour block; TI surfaces come from the layout background
        el = copy.deepcopy(shape)
        place(el)
        remap_colours(el, surface)
        normalise_text(el)
        if is_title(el):
            body = el.find(qn("p:txBody"))
            lst = body.find(qn("a:lstStyle"))
            # on the teal page the title is a small white label above the statement
            label = ' sz="1400"' if surface else ""
            colour = '<a:solidFill><a:schemeClr val="bg1"/></a:solidFill>' if surface else ""
            new = parse_xml(f'<a:lstStyle {nsdecls("a")}><a:lvl1pPr><a:defRPr{label}>{colour}</a:defRPr></a:lvl1pPr></a:lstStyle>')
            if surface:
                body.find(qn("a:bodyPr")).set("anchor", "t")
            if lst is None:
                body.find(qn("a:bodyPr")).addnext(new)
            else:
                body.replace(lst, new)
            for r in body.iter(qn("a:rPr")):
                for attr in ("sz", "b", "i"):
                    r.attrib.pop(attr, None)
                for fill in r.findall(qn("a:solidFill")):
                    r.remove(fill)
        sid, L.next = L.next, L.next + 1
        el.find(".//" + qn("p:cNvPr")).set("id", str(sid))
        L.tree.append(el)
    xml = etree.tostring(L.el).decode()
    if "creationId" in xml:  # copied shapes keep Hoffman's ids; make them unique
        for node in L.el.iter():
            if node.tag.endswith("creationId") and node.get("id"):
                node.set("id", "{" + str(uuid.uuid4()).upper() + "}")
            elif node.tag.endswith("creationId") and node.get("val"):
                node.set("val", str(random.randint(10**8, 2**31 - 1)))


def reorder(ti):
    master = ti.slide_masters[0]
    lst = master._element.find(qn("p:sldLayoutIdLst"))
    by_part = {master.part.related_part(e.get(qn("r:id"))).partname: e for e in lst}
    entries = {l.name: by_part[l.part.partname] for l in ti.slide_layouts}
    for e in list(lst):
        lst.remove(e)
    for name in ORDER:
        if name in entries:
            lst.append(entries.pop(name))
    for e in entries.values():
        lst.append(e)


if __name__ == "__main__":
    ha = Presentation(HA_SRC)
    ti = Presentation(TPL)
    ha_layouts = {l.name: l for l in ha.slide_masters[0].slide_layouts}
    have = {l.name for l in ti.slide_layouts}
    base = next(l for l in ti.slide_layouts if l.name == "Title only")
    for ha_name, name, surface in PORTS:
        if name not in have:
            port(ti, ha_layouts[ha_name], name, surface, base)
    reorder(ti)
    count = len(ti.slide_layouts)
    for shape in ti.slides[0].shapes:  # the "Start here" slide states the layout count
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    r.text = re.sub(r"^ \d+ layouts, including", f" {count} layouts, including", r.text)
    ti.save(OUT)
    print(f"{OUT}: {len(ti.slide_layouts)} layouts")
