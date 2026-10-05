#!/usr/bin/env python3
"""Add the fillable layout set to the TI deck template (template v1.1).

    python3 tools/add_layouts.py [in.pptx] [out.pptx]

Defaults to templates/powerpoint/TI_Deck_Template_v1.pptx in place. Works on TI's
own master: it never rebuilds the master, only adds layouts to it, renames the
kept ones, fills in the theme (named palette, custom colours, object defaults),
adds a hairline default table style and puts a live slide number on content
layouts. Safe to re-run: a layout that already exists by name is left alone.
"""
import random
import re
import sys
import uuid
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.parts.slide import SlideLayoutPart

SRC = sys.argv[1] if len(sys.argv) > 1 else "templates/powerpoint/TI_Deck_Template_v1.pptx"
OUT = sys.argv[2] if len(sys.argv) > 2 else SRC

E = 914400
NS = nsdecls("a", "p", "r")
RED, TEAL, INK, WHITE, CAPTION = "accent1", "accent3", "tx1", "bg1", "bg2"
HAIR, PANEL, TEAL300, WORDMARK = "#CCCCCC", "#F7F7F7", "#4ABED4", "#E8E8E8"
TABLE_STYLE_ID = "{6E25B4C1-2F0A-4C8E-9B1D-7A1C0E0D7101}"
NUMBER_BOX = (7.55, 4.88, 2.2, 0.2)  # page number, right-aligned above the footer rule


def emu(v):
    return int(round(v * E))


def clr(c):
    return f'<a:srgbClr val="{c[1:]}"/>' if c.startswith("#") else f'<a:schemeClr val="{c}"/>'


def xfrm(box):
    x, y, w, h = box
    return f'<a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/><a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>'


# ---------------------------------------------------------------- shape XML

def lvl_xml(n, size, bold, color, caps, algn, bullet, lnspc):
    if bullet:
        attrs = f' marL="{189124 * n}" indent="-189124" algn="{algn}"'
        bu = f'<a:buFont typeface="Arial"/><a:buChar char="{"•" if n == 1 else "–"}"/>'
    else:
        attrs = f' marL="0" indent="0" algn="{algn}"'
        bu = "<a:buNone/>"
    ls = f'<a:lnSpc><a:spcPct val="{lnspc}"/></a:lnSpc>' if lnspc else ""
    r = (f' sz="{int(size * 100)}"' if size else "") + (f' b="{int(bold)}"' if bold is not None else "")
    r += ' cap="all"' if caps else ""
    fill = f"<a:solidFill>{clr(color)}</a:solidFill>" if color else ""
    return f"<a:lvl{n}pPr{attrs}>{ls}{bu}<a:defRPr{r}>{fill}</a:defRPr></a:lvl{n}pPr>"


def ph_xml(sid, name, kind, idx, box, prompt, size=None, bold=None, color=None, caps=False,
           algn="l", anchor="t", bullets=False, lnspc=None, fill=None, line=None, size2=None):
    """A layout placeholder with its own prompt text (shown while editing, never printed)."""
    attrs = (f' type="{kind}"' if kind else "") + (f' idx="{idx}"' if idx is not None else "")
    geom = xfrm(box)
    if fill or line:
        geom += '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        geom += f"<a:solidFill>{clr(fill)}</a:solidFill>" if fill else "<a:noFill/>"
        if line:
            geom += f'<a:ln w="{line[1]}"><a:solidFill>{clr(line[0])}</a:solidFill></a:ln>'
    lst = lvl_xml(1, size, bold, color, caps, algn, bullets, lnspc)
    if bullets:
        lst += lvl_xml(2, size2 or size, bold, color, caps, algn, True, lnspc)
    return (
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(name)}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        f'<p:nvPr><p:ph{attrs} hasCustomPrompt="1"/></p:nvPr></p:nvSpPr>'
        f"<p:spPr>{geom}</p:spPr>"
        f'<p:txBody><a:bodyPr anchor="{anchor}"><a:noAutofit/></a:bodyPr><a:lstStyle>{lst}</a:lstStyle>'
        f'<a:p><a:r><a:rPr lang="en-US" dirty="0"/><a:t>{escape(prompt)}</a:t></a:r></a:p></p:txBody></p:sp>'
    )


def rect_xml(sid, name, box, fill=None, line=None, prst="rect"):
    f = f"<a:solidFill>{clr(fill)}</a:solidFill>" if fill else "<a:noFill/>"
    ln = (f'<a:ln w="{line[1]}"><a:solidFill>{clr(line[0])}</a:solidFill></a:ln>' if line
          else "<a:ln><a:noFill/></a:ln>")
    return (
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(name)}"/><p:cNvSpPr/>'
        f'<p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{xfrm(box)}'
        f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>{f}{ln}</p:spPr></p:sp>'
    )


def line_xml(sid, name, x1, y1, x2, y2, color=HAIR, w=9525):
    box = (min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
    return (
        f'<p:cxnSp {NS}><p:nvCxnSpPr><p:cNvPr id="{sid}" name="{escape(name)}"/><p:cNvCxnSpPr/>'
        f'<p:nvPr userDrawn="1"/></p:nvCxnSpPr><p:spPr>{xfrm(box)}'
        f'<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
        f'<a:ln w="{w}"><a:solidFill>{clr(color)}</a:solidFill></a:ln></p:spPr></p:cxnSp>'
    )


def text_xml(sid, name, box, text, size, color, bold=False, algn="l", anchor="t", slidenum=False):
    rpr = f'<a:rPr lang="en-US" sz="{int(size * 100)}" b="{int(bold)}" dirty="0"><a:solidFill>{clr(color)}</a:solidFill></a:rPr>'
    if slidenum:
        run = f'<a:fld id="{{{str(uuid.uuid4()).upper()}}}" type="slidenum">{rpr}<a:t>‹#›</a:t></a:fld>'
    else:
        run = f"<a:r>{rpr}<a:t>{escape(text)}</a:t></a:r>"
    return (
        f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{sid}" name="{escape(name)}"/><p:cNvSpPr txBox="1"/>'
        f'<p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr>{xfrm(box)}'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}">'
        f'<a:noAutofit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="{algn}"/>{run}</a:p></p:txBody></p:sp>'
    )


# ---------------------------------------------------------------- layout plumbing

class Layout:
    """Wraps one slideLayout element: add shapes with fresh ids, drop shapes, set backgrounds."""

    def __init__(self, layout):
        self.layout = layout
        self.el = layout._element
        self.tree = self.el.find(qn("p:cSld")).find(qn("p:spTree"))
        self.next = max([int(e.get("id")) for e in self.el.iter(qn("p:cNvPr"))] + [1]) + 1

    def add(self, maker, *args, **kw):
        sid, self.next = self.next, self.next + 1
        self.tree.append(parse_xml(maker(sid, *args, **kw)))

    def ph(self, *args, **kw):
        self.add(ph_xml, *args, **kw)

    def drop(self, pred):
        for shape in list(self.tree):
            if shape.tag in (qn("p:sp"), qn("p:pic"), qn("p:cxnSp")) and pred(shape):
                self.tree.remove(shape)

    def drop_title(self):
        self.drop(lambda s: _ph_type(s) == "title")

    def drop_picture(self):
        self.drop(lambda s: s.tag == qn("p:pic"))

    def background(self, color, hide_master=True):
        csld = self.el.find(qn("p:cSld"))
        old = csld.find(qn("p:bg"))
        if old is not None:
            csld.remove(old)
        csld.insert(0, parse_xml(f'<p:bg {NS}><p:bgPr><a:solidFill>{clr(color)}</a:solidFill>'
                                 f"<a:effectLst/></p:bgPr></p:bg>"))
        if hide_master:
            self.el.set("showMasterSp", "0")

    def prompt(self, ph_type, text, idx=None):
        """Give an inherited placeholder (title, body…) a custom prompt."""
        for s in self.tree.iter(qn("p:sp")):
            ph = s.find(".//" + qn("p:ph"))
            if ph is None or (ph.get("type") or "obj") != ph_type:
                continue
            if idx is not None and ph.get("idx") != str(idx):
                continue
            ph.set("hasCustomPrompt", "1")
            body = s.find(qn("p:txBody"))
            for p in body.findall(qn("a:p")):
                body.remove(p)
            body.append(parse_xml(f'<a:p {NS}><a:r><a:rPr lang="en-US" dirty="0"/>'
                                  f"<a:t>{escape(text)}</a:t></a:r></a:p>"))
            return s

    def geometry(self, ph_type, box, idx=None):
        s = self.prompt_target(ph_type, idx)
        sppr = s.find(qn("p:spPr"))
        old = sppr.find(qn("a:xfrm"))
        if old is not None:
            sppr.remove(old)
        sppr.insert(0, parse_xml(f"<a:xfrm {NS}>" + xfrm(box)[len("<a:xfrm>"):]))
        return s

    def prompt_target(self, ph_type, idx=None):
        for s in self.tree.iter(qn("p:sp")):
            ph = s.find(".//" + qn("p:ph"))
            if ph is not None and (ph.get("type") or "obj") == ph_type and (idx is None or ph.get("idx") == str(idx)):
                return s
        raise KeyError(ph_type)

    def style(self, ph_type, idx=None, **kw):
        """Replace the list style of an inherited placeholder (size, colour, anchor…)."""
        s = self.prompt_target(ph_type, idx)
        body = s.find(qn("p:txBody"))
        anchor = kw.pop("anchor", None)
        bullets = kw.pop("bullets", False)
        size2 = kw.pop("size2", None)
        args = dict(size=None, bold=None, color=None, caps=False, algn="l", lnspc=None) | kw
        lst = lvl_xml(1, args["size"], args["bold"], args["color"], args["caps"], args["algn"], bullets, args["lnspc"])
        if bullets:
            lst += lvl_xml(2, size2 or args["size"], args["bold"], args["color"], args["caps"], args["algn"], True, args["lnspc"])
        new = parse_xml(f"<a:lstStyle {NS}>{lst}</a:lstStyle>")
        old = body.find(qn("a:lstStyle"))
        if old is None:
            body.find(qn("a:bodyPr")).addnext(new)
        else:
            body.replace(old, new)
        if anchor:
            body.find(qn("a:bodyPr")).set("anchor", anchor)

    def slide_number(self):
        if any(c.get("name") == "Slide number" for c in self.tree.iter(qn("p:cNvPr"))):
            return  # inherited from the layout this one was cloned from
        self.add(text_xml, "Slide number", NUMBER_BOX, "", 9, CAPTION, algn="r", anchor="b", slidenum=True)


def _ph_type(shape):
    ph = shape.find(".//" + qn("p:ph"))
    return None if ph is None else (ph.get("type") or "obj")


def clone_layout(prs, src, name):
    """Copy a layout part (relationships and all) and register it on the master."""
    master = prs.slide_masters[0]
    xml = etree.tostring(src._element).decode()
    xml = re.sub(r'(creationId [^>]*?val=")\d+', lambda m: m.group(1) + str(random.randint(10**8, 2**31 - 1)), xml)
    xml = re.sub(r'(creationId [^>]*?id=")\{[0-9A-Fa-f-]+\}', lambda m: m.group(1) + "{" + str(uuid.uuid4()).upper() + "}", xml)
    el = parse_xml(xml)
    pkg = prs.part.package
    used = {p.partname for p in pkg.iter_parts()}
    n = 1
    while PackURI(f"/ppt/slideLayouts/slideLayout{n}.xml") in used:
        n += 1
    part = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{n}.xml"), src.part.content_type, pkg, el)
    rmap = {rid: part.relate_to(rel.target_part, rel.reltype)
            for rid, rel in src.part.rels.items() if not rel.is_external}
    for node in el.iter():
        for attr in (qn("r:embed"), qn("r:id"), qn("r:link")):
            if node.get(attr) in rmap:
                node.set(attr, rmap[node.get(attr)])
    rid = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    lst = master._element.find(qn("p:sldLayoutIdLst"))
    ids = [int(e.get("id")) for e in lst] + [int(e.get("id")) for e in prs.part._element.find(qn("p:sldMasterIdLst"))]
    entry = etree.SubElement(lst, qn("p:sldLayoutId"))
    entry.set("id", str(max(ids) + 1))
    entry.set(qn("r:id"), rid)
    el.find(qn("p:cSld")).set("name", name)
    if "type" in el.attrib:
        del el.attrib["type"]
    el.set("preserve", "1")
    return next(l for l in prs.slide_layouts if l.part is part)


# ---------------------------------------------------------------- the layout set

def build(prs):
    by_name = {l.name: l for l in prs.slide_layouts}
    names = set(by_name)

    def base(old_name, new_name):
        """Rename a kept layout (first run) or find it again (re-run)."""
        lay = by_name.get(old_name) or by_name[new_name]
        lay._element.find(qn("p:cSld")).set("name", new_name)
        return lay

    cover = base("Title Slide", "Cover · L01")
    closing = base("5_Title Slide", "Closing · L30")
    body = base("Title and Content", "Title + body")
    two = base("Two Content", "Two columns")
    title_only = base("Title Only", "Title only")
    blank = base("Blank", "Blank")
    if "Statement · L08" in names:  # already upgraded
        return

    title_prompt = "Title states the point: a claim, not a topic (two lines at most)"

    # Cover: eyebrow, title, subtitle on the grey-circuit background
    L = Layout(cover)
    L.geometry("ctrTitle", (0.25, 1.7, 6.4, 1.3))
    L.style("ctrTitle", size=30, bold=True, color=RED, anchor="b", lnspc=85000)
    L.prompt("ctrTitle", "The one idea this deck exists to land")
    L.geometry("subTitle", (0.25, 3.05, 6.4, 0.6), idx=1)
    L.style("subTitle", idx=1, size=14, color=INK)
    L.prompt("subTitle", "Who this is for and what they get, in one line", idx=1)
    L.ph("Eyebrow", "body", 13, (0.25, 1.35, 6.0, 0.3), "CLIENT · ENGAGEMENT · MONTH YEAR",
         size=10, bold=True, color=TEAL, caps=True)

    # Closing: imperative headline + contact on the light-facets background
    L = Layout(closing)
    L.geometry("ctrTitle", (0.25, 1.8, 7.5, 1.0))
    L.style("ctrTitle", size=30, bold=True, color=RED, anchor="b", lnspc=85000)
    L.prompt("ctrTitle", "Close with a verb, six words at most")
    L.geometry("subTitle", (0.25, 3.0, 7.5, 0.5), idx=1)
    L.style("subTitle", idx=1, size=14, color=INK)
    L.prompt("subTitle", "Name · The Hoffman Agency · email", idx=1)

    # Title + body and Two columns: body at 14/12pt (the master's 18pt is too big for TI pages)
    L = Layout(body)
    L.prompt("title", title_prompt)
    L.style("obj", idx=1, size=14, size2=12, bullets=True)
    L.prompt("obj", "Points that prove the title. Short bullets, never paragraphs.", idx=1)
    L.slide_number()

    L = Layout(two)
    L.prompt("title", title_prompt)
    for idx, x in ((1, 0.25), (2, 5.08)):
        L.geometry("obj", (x, 1.4, 4.55, 3.35), idx=idx)
        L.style("obj", idx=idx, size=14, size2=12, bullets=True)
        L.prompt("obj", "Points for this column", idx=idx)
    L.ph("Left heading", "body", 13, (0.25, 0.92, 4.55, 0.4), "Column heading", size=16, bold=True, color=TEAL)
    L.ph("Right heading", "body", 14, (5.08, 0.92, 4.55, 0.4), "Column heading", size=16, bold=True, color=TEAL)
    L.slide_number()

    L = Layout(title_only)
    L.prompt("title", title_prompt)
    L.slide_number()

    # Section divider (red) and appendix divider (teal): no signature on a coloured surface
    for name, colour in (("Section divider · L02", RED), ("Appendix divider · L02 teal", TEAL)):
        L = Layout(clone_layout(prs, blank, name))
        L.drop_picture()
        L.background(colour)
        L.ph("Part number", "body", 13, (0.25, 1.5, 3.0, 1.2), "01", size=72, bold=True, color=WHITE, anchor="b")
        L.ph("Title", "title", None, (0.25, 3.0, 9.0, 1.0), "Section name, one or two words",
             size=36, bold=True, color=WHITE, lnspc=85000)

    # Agenda: five numbered rows, grey AGENDA wordmark
    L = Layout(clone_layout(prs, title_only, "Agenda · L03"))
    L.prompt("title", "Agenda")
    L.add(text_xml, "AGENDA wordmark", (6.6, 1.1, 3.15, 0.9), "AGENDA", 40, WORDMARK, bold=True, algn="r")
    for i in range(5):
        y = 1.0 + i * 0.76
        L.ph(f"Number {i + 1}", "body", 13 + 2 * i, (0.25, y, 0.9, 0.6), f"0{i + 1}",
             size=28, bold=True, color=TEAL, anchor="ctr")
        L.ph(f"Item {i + 1}", "body", 14 + 2 * i, (1.25, y + 0.08, 5.4, 0.45), "Part name, four words at most",
             size=16, color=INK, anchor="ctr")
        L.add(line_xml, f"Rule {i + 1}", 0.25, y + 0.68, 6.85, y + 0.68)

    # Statement + support: the claim is the title, so the outline view still reads as an argument
    L = Layout(clone_layout(prs, title_only, "Statement · L08"))
    L.drop_title()
    L.ph("Claim", "title", None, (0.25, 0.9, 8.6, 2.0), "One claim in a sentence. Colour one phrase red.",
         size=34, bold=True, color=INK, lnspc=90000)
    L.ph("Support", "body", 13, (0.25, 3.1, 6.2, 1.2), "Support: at most 30 words that extend the claim",
         size=14, color=INK)
    L.ph("Evidence", "body", 14, (7.0, 3.05, 2.75, 1.75), "Evidence slot: one stat, one quote or one chart. Not three.",
         size=11, color=INK, line=(TEAL300, 19050))
    L.slide_number()

    # Three columns: icon, teal heading, two bullets per card
    L = Layout(clone_layout(prs, title_only, "Three columns · L12"))
    L.prompt("title", title_prompt)
    for i, x in enumerate((0.25, 3.52, 6.78)):
        L.add(rect_xml, f"Card {i + 1}", (x, 0.9, 2.97, 3.9), line=(HAIR, 9525))
        L.ph(f"Icon {i + 1}", "pic", 13 + 3 * i, (x + 0.25, 1.15, 0.7, 0.7), "Icon", size=9, color=CAPTION, algn="ctr",
             anchor="ctr", fill=PANEL)
        L.ph(f"Heading {i + 1}", "body", 14 + 3 * i, (x + 0.25, 2.0, 2.47, 0.6), "Heading, seven words at most",
             size=16, bold=True, color=TEAL)
        L.ph(f"Body {i + 1}", "body", 15 + 3 * i, (x + 0.25, 2.7, 2.47, 1.9), "Two short bullets",
             size=12, color=INK, bullets=True)
    L.slide_number()

    # Stat strip: one hero (red) and two context figures, with a source line
    L = Layout(clone_layout(prs, title_only, "Stat strip · L19"))
    L.prompt("title", title_prompt)
    for i, (x, size, colour, prompt) in enumerate(((0.55, 60, RED, "Hero"), (3.72, 44, INK, "Stat"),
                                                     (6.88, 44, INK, "Stat"))):
        L.ph(f"Number {i + 1}", "body", 13 + 2 * i, (x, 1.7, 2.57, 1.3), prompt, size=size, bold=True, color=colour,
             anchor="b")
        L.ph(f"Label {i + 1}", "body", 14 + 2 * i, (x, 3.1, 2.57, 0.6), "Label, four words at most", size=12, color=INK)
    for x in (3.42, 6.58):
        L.add(line_xml, "Divider", x, 1.7, x, 4.1)
    L.ph("Source", "body", 19, (0.25, 4.5, 7.2, 0.3), "Source: who, what, when — or [REAL DATA · what · who supplies it]",
         size=9, color=CAPTION, anchor="b")
    L.slide_number()

    # Chart + takeaway
    L = Layout(clone_layout(prs, title_only, "Chart + takeaway"))
    L.prompt("title", title_prompt)
    L.ph("Chart", "chart", 13, (0.25, 0.9, 6.3, 3.55), "Chart: TI in red, everyone else teal and grey")
    L.ph("Takeaway", "body", 14, (6.8, 0.9, 2.95, 3.55), "What the chart proves, in 30 words at most",
         size=14, color=INK)
    L.ph("Source", "body", 15, (0.25, 4.5, 7.2, 0.3), "Source: who, what, when — or [REAL DATA · what · who supplies it]",
         size=9, color=CAPTION, anchor="b")
    L.slide_number()

    # Table (the one dense page)
    L = Layout(clone_layout(prs, title_only, "Table · L24"))
    L.prompt("title", title_prompt)
    L.ph("Table", "tbl", 13, (0.25, 0.9, 9.5, 3.5), "Table: hairline rows, grey header. One or two per deck.")
    L.ph("Note", "body", 14, (0.25, 4.5, 7.2, 0.3), "Source or note", size=9, color=CAPTION, anchor="b")
    L.slide_number()

    # Image + text: photo bleeds the left half
    L = Layout(clone_layout(prs, title_only, "Image + text · L28"))
    L.drop_title()
    L.ph("Photo", "pic", 13, (0.0, 0.0, 5.0, 5.625), "Photo: TI imagery only; people shown are Taiwanese",
         size=10, color=CAPTION, algn="ctr", anchor="ctr", fill=PANEL)
    L.ph("Eyebrow", "body", 14, (5.4, 0.9, 4.3, 0.3), "EYEBROW", size=10, bold=True, color=TEAL, caps=True)
    L.ph("Title", "title", None, (5.4, 1.25, 4.3, 1.6), "Headline that states the point",
         size=26, bold=True, color=RED, lnspc=85000)
    L.ph("Body", "body", 15, (5.4, 2.95, 4.3, 1.4), "Body, 30 words at most", size=12, color=INK)
    L.ph("Call to action", "body", 16, (5.4, 4.4, 4.3, 0.3), "Next step →", size=11, bold=True, color=TEAL)
    L.slide_number()

    # Timeline: five nodes on a hairline
    L = Layout(clone_layout(prs, title_only, "Timeline · L20"))
    L.prompt("title", title_prompt)
    L.add(line_xml, "Timeline rule", 0.25, 2.4, 9.75, 2.4)
    for i in range(5):
        cx = 1.2 + i * 1.9
        L.add(rect_xml, f"Node {i + 1}", (cx - 0.09, 2.31, 0.18, 0.18), fill=TEAL, prst="ellipse")
        L.ph(f"When {i + 1}", "body", 13 + 3 * i, (cx - 0.95, 1.7, 1.9, 0.4), "WHEN", size=10, bold=True,
             color=TEAL, caps=True, algn="ctr", anchor="b")
        L.ph(f"Step {i + 1}", "body", 14 + 3 * i, (cx - 0.95, 2.65, 1.9, 0.4), "Step", size=16, bold=True,
             color=INK, algn="ctr")
        L.ph(f"Note {i + 1}", "body", 15 + 3 * i, (cx - 0.85, 3.1, 1.7, 1.2), "One line, 15 words at most",
             size=11, color=INK, algn="ctr")
    L.slide_number()

    # Team: four people
    L = Layout(clone_layout(prs, title_only, "Team · L15"))
    L.prompt("title", title_prompt)
    for i in range(4):
        x = 0.25 + i * 2.45
        L.ph(f"Photo {i + 1}", "pic", 13 + 3 * i, (x, 0.95, 2.15, 2.1), "Photo", size=10, color=CAPTION,
             algn="ctr", anchor="ctr", fill=PANEL)
        L.ph(f"Name {i + 1}", "body", 14 + 3 * i, (x, 3.12, 2.15, 0.35), "Name", size=14, bold=True, color=INK)
        L.ph(f"Role {i + 1}", "body", 15 + 3 * i, (x, 3.47, 2.15, 1.0), "Role on the account, and the most credible fact",
             size=11, color=INK)
    L.slide_number()

    # Pull quote on teal
    L = Layout(clone_layout(prs, blank, "Pull quote · L29"))
    L.drop_picture()
    L.background(TEAL)
    L.add(text_xml, "Quote mark", (0.45, 0.75, 1.0, 1.2), "“", 96, WHITE, bold=True)
    L.ph("Quote", "body", 13, (0.45, 1.7, 8.5, 2.0), "A real quote, 20 words at most", size=28, color=WHITE)
    L.ph("Attribution", "body", 14, (0.45, 3.9, 8.0, 0.5), "Who said it, and in what role", size=12, color=WHITE)

    # menu order
    order = ["Cover · L01", "Section divider · L02", "Appendix divider · L02 teal", "Agenda · L03", "Title + body",
             "Two columns", "Statement · L08", "Three columns · L12", "Stat strip · L19", "Chart + takeaway",
             "Table · L24", "Image + text · L28", "Timeline · L20", "Team · L15", "Pull quote · L29", "Title only",
             "Closing · L30", "Blank"]
    master = prs.slide_masters[0]
    lst = master._element.find(qn("p:sldLayoutIdLst"))
    rid_of = {master.part.related_part(e.get(qn("r:id"))).partname: e for e in lst}
    entries = {l.name: rid_of[l.part.partname] for l in prs.slide_layouts}
    for e in list(lst):
        lst.remove(e)
    for name in order:
        lst.append(entries.pop(name))
    for e in entries.values():  # anything unexpected stays, at the end
        lst.append(e)


# ---------------------------------------------------------------- theme, tables, example slides

CUSTOM_COLOURS = [("TI red · title, rule, one hero", "CC0000"), ("Teal · icons, lead-ins", "117788"),
                  ("Teal 500 · chart", "2790A5"), ("Teal 400 · chart", "32B4CE"), ("Teal 300 · callout fill", "4ABED4"),
                  ("Grey dark · secondary text", "404040"), ("Grey mid", "7F7F7F"), ("Grey light · chart", "A4A4A4"),
                  ("Hairline", "CCCCCC"), ("Panel · grey 100", "F7F7F7")]

OBJECT_DEFAULTS = (
    "<a:objectDefaults>"
    # new shapes: teal fill, white Arial 12pt (teal is the working accent; red stays for the one hero)
    '<a:spDef><a:spPr><a:solidFill><a:schemeClr val="accent3"/></a:solidFill><a:ln><a:noFill/></a:ln></a:spPr>'
    '<a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle><a:defPPr algn="ctr"><a:defRPr sz="1200" dirty="0">'
    '<a:solidFill><a:schemeClr val="bg1"/></a:solidFill></a:defRPr></a:defPPr></a:lstStyle>'
    '<a:style><a:lnRef idx="0"><a:schemeClr val="accent3"/></a:lnRef><a:fillRef idx="0"><a:schemeClr val="accent3"/>'
    '</a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent3"/></a:effectRef><a:fontRef idx="minor">'
    '<a:schemeClr val="bg1"/></a:fontRef></a:style></a:spDef>'
    # new lines and arrows: teal 1.5pt
    '<a:lnDef><a:spPr><a:ln w="19050"><a:solidFill><a:schemeClr val="accent3"/></a:solidFill></a:ln></a:spPr>'
    '<a:bodyPr/><a:lstStyle/><a:style><a:lnRef idx="1"><a:schemeClr val="accent3"/></a:lnRef><a:fillRef idx="0">'
    '<a:schemeClr val="accent3"/></a:fillRef><a:effectRef idx="0"><a:schemeClr val="accent3"/></a:effectRef>'
    '<a:fontRef idx="minor"><a:schemeClr val="tx1"/></a:fontRef></a:style></a:lnDef>'
    # new text boxes: Arial 12pt black, the body floor
    '<a:txDef><a:spPr><a:noFill/></a:spPr><a:bodyPr wrap="square" rtlCol="0"><a:spAutoFit/></a:bodyPr>'
    '<a:lstStyle><a:defPPr><a:defRPr sz="1200" dirty="0"/></a:defPPr></a:lstStyle></a:txDef>'
    "</a:objectDefaults>"
)


def patch_theme(prs):
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    x = part.blob.decode("utf-8")
    x = re.sub(r'(<a:theme [^>]*name=")[^"]*"', r'\1Texas Instruments"', x, count=1)
    x = re.sub(r'<a:clrScheme name="[^"]*"', '<a:clrScheme name="Texas Instruments"', x, count=1)
    x = re.sub(r'<a:fontScheme name="[^"]*"', '<a:fontScheme name="Texas Instruments · Arial"', x, count=1)
    x = re.sub(r"<a:objectDefaults/>|<a:objectDefaults>.*?</a:objectDefaults>", OBJECT_DEFAULTS, x, count=1, flags=re.S)
    cust = "".join(f'<a:custClr name="{escape(n)}"><a:srgbClr val="{v}"/></a:custClr>' for n, v in CUSTOM_COLOURS)
    x = re.sub(r"<a:custClrLst>.*?</a:custClrLst>", "", x, flags=re.S)
    x = x.replace("</a:extraClrSchemeLst>", f"</a:extraClrSchemeLst><a:custClrLst>{cust}</a:custClrLst>", 1)
    part._blob = x.encode("utf-8")


def patch_table_styles(prs):
    part = prs.part.part_related_by(RT.TABLE_STYLES)
    x = part.blob.decode("utf-8")
    hair = '<a:ln w="9525"><a:solidFill><a:srgbClr val="CCCCCC"/></a:solidFill></a:ln>'
    none = '<a:ln w="9525"><a:noFill/></a:ln>'
    txt = '<a:fontRef idx="minor"><a:prstClr val="black"/></a:fontRef><a:schemeClr val="tx1"/>'
    style = (
        f'<a:tblStyle styleId="{TABLE_STYLE_ID}" styleName="TI hairline">'
        f"<a:wholeTbl><a:tcTxStyle>{txt}</a:tcTxStyle><a:tcStyle><a:tcBdr>"
        f"<a:left>{none}</a:left><a:right>{none}</a:right><a:top>{hair}</a:top><a:bottom>{hair}</a:bottom>"
        f"<a:insideH>{hair}</a:insideH><a:insideV>{none}</a:insideV></a:tcBdr><a:fill><a:noFill/></a:fill>"
        f"</a:tcStyle></a:wholeTbl>"
        f'<a:firstRow><a:tcTxStyle b="on">{txt}</a:tcTxStyle><a:tcStyle><a:tcBdr><a:bottom>'
        f'<a:ln w="12700"><a:solidFill><a:srgbClr val="CCCCCC"/></a:solidFill></a:ln></a:bottom></a:tcBdr>'
        f'<a:fill><a:solidFill><a:srgbClr val="F7F7F7"/></a:solidFill></a:fill></a:tcStyle></a:firstRow>'
        f"</a:tblStyle>"
    )
    if TABLE_STYLE_ID not in x:
        x = re.sub(r"(<a:tblStyleLst[^>]*?)\s*/>", r"\1></a:tblStyleLst>", x, count=1)
        x = x.replace("</a:tblStyleLst>", style + "</a:tblStyleLst>", 1)
    x = re.sub(r'(<a:tblStyleLst[^>]*\bdef=")[^"]*"', rf'\g<1>{TABLE_STYLE_ID}"', x, count=1)
    part._blob = x.encode("utf-8")


def tidy_examples(prs):
    """Example slides: drop typed page numbers (layouts now number slides) and decimal EMUs."""
    for slide in prs.slides:
        for node in slide._element.iter():
            for attr in ("x", "y", "cx", "cy"):
                v = node.get(attr)
                if v and re.fullmatch(r"-?\d+\.\d+", v):
                    node.set(attr, str(int(round(float(v)))))
        for shape in list(slide.shapes):
            if (shape.has_text_frame and shape.text_frame.text.strip().isdigit()
                    and abs(shape.left - emu(7.55)) < emu(0.05) and abs(shape.top - emu(4.9)) < emu(0.05)):
                shape._element.getparent().remove(shape._element)
    start = prs.slides[0]
    for shape in start.shapes:
        if not shape.has_text_frame:
            continue
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                if r.text == "Pick a layout, duplicate it, replace the text.":
                    r.text = "New Slide → pick a layout, then type into the boxes."
                elif r.text.startswith(" Every slide carries its layout code"):
                    r.text = (" 18 layouts, including red section dividers and a teal appendix divider. The grey prompt "
                              "text says what goes where and never prints; the example slides show finished versions.")
                elif r.text.startswith("Texas Instruments design system v1 ·"):
                    r.text = r.text.replace("design system v1 ·", "design system v1.1 ·")


if __name__ == "__main__":
    prs = Presentation(SRC)
    first_run = "Statement · L08" not in {l.name for l in prs.slide_layouts}
    build(prs)
    patch_theme(prs)
    patch_table_styles(prs)
    if first_run:
        tidy_examples(prs)
    prs.save(OUT)
    print(f"{OUT}: {len(prs.slide_layouts)} layouts, {len(prs.slides)} slides")
    for l in prs.slide_layouts:
        print("  ", l.name)
